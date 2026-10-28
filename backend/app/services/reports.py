"""Report drafting, four-eyes validation and marked exports (PDF, DOCX)."""

from __future__ import annotations

import hashlib
import io
import json
import secrets
import uuid

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ApiError, bad_request, conflict
from app.core.security import utcnow
from app.models import (
    Case,
    ClassificationLevel,
    ExportRecord,
    Report,
    Source,
    Statement,
    User,
    Validation,
)
from app.schemas import api as s
from app.services import audit
from app.services.notifications import notify

REPORT_TITLES = {
    "intelligence_note": ("Note de renseignement", "Intelligence note"),
    "daily_summary": ("Synthèse quotidienne", "Daily summary"),
    "situation_report": ("Rapport de situation", "Situation report"),
    "threat_assessment": ("Évaluation de menace", "Threat assessment"),
    "incident_report": ("Rapport d'incident", "Incident report"),
    "analytical_chronology": ("Chronologie analytique", "Analytical chronology"),
    "strategic_watch_note": ("Note de veille stratégique", "Strategic watch note"),
    "executive_briefing": ("Briefing exécutif", "Executive briefing"),
    "case_closure_report": ("Rapport de clôture", "Case closure report"),
}
KIND_LABELS = {
    "confirmed_fact": "Fait confirmé",
    "unverified_information": "Information non vérifiée",
    "claim": "Déclaration",
    "working_hypothesis": "Hypothèse de travail",
    "alternative_hypothesis": "Hypothèse alternative",
}
HUMAN_REVIEW_NOTICE = (
    "Document d'aide à la décision. Il ne constitue pas une preuve et ne justifie à lui seul aucune mesure coercitive."
)


def content_hash(report: Report) -> str:
    payload = {k: getattr(report, k) for k in (
        "title", "report_type", "classification_rank", "dissemination", "summary", "analysis",
        "uncertainties", "recommendations")}
    payload["source_ids"] = sorted(map(str, report.source_ids))
    payload["statement_ids"] = sorted(map(str, report.statement_ids))
    return hashlib.sha256(json.dumps(payload, sort_keys=True).encode()).hexdigest()


def _validate_refs(db: Session, case: Case, source_ids, statement_ids) -> None:
    for sid in source_ids:
        src = db.get(Source, sid)
        if src is None or src.case_id != case.id:
            raise bad_request("source_not_in_case")
    for stid in statement_ids:
        st = db.get(Statement, stid)
        if st is None or st.case_id != case.id:
            raise bad_request("statement_not_in_case")


def _prefill(db: Session, case: Case) -> dict:
    statements = db.scalars(select(Statement).where(Statement.case_id == case.id).order_by(Statement.created_at)).all()
    sections = []
    for kind, label in KIND_LABELS.items():
        items = [st for st in statements if st.kind == kind]
        if items:
            sections.append(f"{label}s :\n" + "\n".join(f"[{st.id.hex[:8]}] {st.text}" for st in items))
    uncertain = [st for st in statements if st.kind != "confirmed_fact"]
    return {
        "summary": case.executive_summary,
        "analysis": "\n\n".join(sections),
        "uncertainties": (
            f"{len(uncertain)} élément(s) non confirmé(s) ou hypothétique(s) figurent dans le dossier."
            if uncertain else ""
        ),
        "statement_ids": [str(st.id) for st in statements],
        "source_ids": sorted({str(st.source_id) for st in statements if st.source_id}),
    }


def create_report(db: Session, principal, case: Case, data: s.ReportIn) -> Report:
    rank = data.classification_rank if data.classification_rank is not None else case.classification_rank
    if rank < case.classification_rank:
        raise bad_request("report_below_case_classification")
    if rank > principal.user.clearance_rank:
        raise ApiError(403, "classification_above_clearance")
    _validate_refs(db, case, data.source_ids, data.statement_ids)
    fields = data.model_dump(exclude={"prefill_from_case", "classification_rank"})
    fields["source_ids"] = [str(x) for x in data.source_ids]
    fields["statement_ids"] = [str(x) for x in data.statement_ids]
    if data.prefill_from_case:
        pre = _prefill(db, case)
        for key in ("summary", "analysis", "uncertainties"):
            fields[key] = fields[key] or pre[key]
        fields["source_ids"] = fields["source_ids"] or pre["source_ids"]
        fields["statement_ids"] = fields["statement_ids"] or pre["statement_ids"]
    report = Report(
        reference=f"RPT-{case.reference}-{secrets.token_hex(2).upper()}", case_id=case.id,
        classification_rank=rank, author_id=principal.id, **fields,
    )
    db.add(report)
    db.flush()
    audit.record(db, principal.audit, "report.create", "report", "success", "report", report.id,
                 {"case_id": str(case.id), "type": data.report_type}, commit=False)
    db.commit()
    return report


def update_report(db: Session, principal, case: Case, report: Report, data: s.ReportUpdateIn) -> Report:
    if report.status not in {"draft", "rejected"}:
        raise conflict("report_locked", "Only drafts or rejected reports can be edited")
    if report.author_id != principal.id:
        raise ApiError(403, "only_author_can_edit")
    if data.version != report.version:
        raise conflict("version_conflict")
    changes = data.model_dump(exclude_unset=True, exclude={"version"})
    _validate_refs(db, case, changes.get("source_ids", []), changes.get("statement_ids", []))
    for key in ("source_ids", "statement_ids"):
        if key in changes:
            changes[key] = [str(x) for x in changes[key]]
    for key, value in changes.items():
        setattr(report, key, value)
    report.status = "draft"
    audit.record(db, principal.audit, "report.update", "report", "success", "report", report.id,
                 {"fields": sorted(changes)}, commit=False)
    db.commit()
    return report


def submit_report(db: Session, principal, case: Case, report: Report) -> Report:
    if report.status not in {"draft", "rejected"}:
        raise conflict("invalid_transition")
    if report.author_id != principal.id:
        raise ApiError(403, "only_author_can_submit")
    if not report.source_ids:
        raise bad_request("report_requires_sources")
    report.status = "submitted"
    report.content_sha256 = content_hash(report)
    validators = _potential_validators(db, case, exclude=principal.id)
    notify(db, validators, "report_validation", "report", report.id)
    audit.record(db, principal.audit, "report.submit", "report", "success", "report", report.id,
                 {"sha256": report.content_sha256}, commit=False)
    db.commit()
    return report


def _potential_validators(db: Session, case: Case, exclude) -> list:
    from app.authz.policy import decide_case

    users = db.scalars(select(User).where(User.is_active.is_(True))).all()
    return [u.id for u in users if u.id != exclude and "report:validate" in u.permission_codes
            and decide_case(db, u, case, "read").allowed]


def decide_report(db: Session, principal, report: Report, decision: str, comment: str) -> Report:
    if report.status != "submitted":
        raise conflict("invalid_transition")
    if report.author_id == principal.id:
        raise ApiError(403, "four_eyes_required", "The author cannot validate their own report")
    if content_hash(report) != report.content_sha256:
        raise conflict("content_changed_since_submission")
    if decision == "rejected" and len(comment) < 3:
        raise bad_request("rejection_requires_comment")
    db.add(Validation(object_type="report", object_id=report.id, decision=decision, comment=comment,
                      validator_id=principal.id, content_sha256=report.content_sha256, decided_at=utcnow()))
    report.status = decision
    notify(db, [report.author_id], "report_decision", "report", report.id)
    audit.record(db, principal.audit, f"report.{decision}", "validation", "success", "report", report.id,
                 {"sha256": report.content_sha256}, commit=False)
    db.commit()
    return report


def export_report(db: Session, principal, case: Case, report: Report, fmt: str) -> tuple[bytes, str, str]:
    if report.status not in {"approved", "disseminated"}:
        raise conflict("report_not_approved", "Only validated reports can be exported")
    if content_hash(report) != report.content_sha256:
        raise conflict("content_changed_since_validation")
    level = db.scalar(select(ClassificationLevel).where(ClassificationLevel.rank == report.classification_rank))
    if level is None:
        raise ApiError(500, "classification_missing")
    if level.export_requires_step_up and not principal.step_up_valid:
        audit.record(db, principal.audit, "report.export", "export", "denied", "report", report.id,
                     {"reason": "step_up_required"})
        raise ApiError(403, "step_up_required")
    trace = f"EXP-{utcnow():%Y%m%d}-{secrets.token_hex(4).upper()}"
    doc = _document_model(db, case, report, level, principal.user, trace)
    data = render_pdf(doc) if fmt == "pdf" else render_docx(doc)
    digest = hashlib.sha256(data).hexdigest()
    db.add(ExportRecord(report_id=report.id, user_id=principal.id, format=fmt, trace_reference=trace,
                        sha256=digest, created_at=utcnow()))
    report.status = "disseminated"
    audit.record(db, principal.audit, "report.export", "export", "success", "report", report.id,
                 {"format": fmt, "trace": trace, "sha256": digest}, commit=False)
    db.commit()
    mime = "application/pdf" if fmt == "pdf" else \
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return data, mime, f"{report.reference}.{fmt}"


def _document_model(db: Session, case: Case, report: Report, level: ClassificationLevel, user: User, trace: str) -> dict:
    sources = [db.get(Source, uuid.UUID(str(sid))) for sid in report.source_ids]
    statements = [db.get(Statement, uuid.UUID(str(sid))) for sid in report.statement_ids]
    validation = db.scalar(select(Validation).where(Validation.object_id == report.id,
                                                    Validation.decision == "approved")
                           .order_by(Validation.decided_at.desc()))
    validator = db.get(User, validation.validator_id) if validation else None
    author = db.get(User, report.author_id)
    fr, en = REPORT_TITLES[report.report_type]
    return {
        "marking": level.marking,
        "dissemination": report.dissemination,
        "type": f"{fr} / {en}",
        "title": report.title,
        "reference": report.reference,
        "case_reference": case.reference,
        "date": utcnow().strftime("%Y-%m-%d %H:%M UTC"),
        "author": author.professional_id if author else "",
        "validator": validator.professional_id if validator else "",
        "trace": trace,
        "exported_by": user.professional_id,
        "content_sha256": report.content_sha256,
        "sections": [
            ("Synthèse / Summary", report.summary),
            ("Analyse / Analysis", report.analysis),
            ("Incertitudes / Uncertainties", report.uncertainties),
            ("Recommandations / Recommendations", report.recommendations),
        ],
        "statements": [(KIND_LABELS.get(st.kind, st.kind), st.text, st.id.hex[:8]) for st in statements if st],
        "sources": [
            f"[{src.id.hex[:8]}] {src.title} ({src.kind}, fiabilité {src.reliability}"
            + (f", {src.url}" if src.url else "") + ")"
            for src in sources if src
        ],
        "notice": HUMAN_REVIEW_NOTICE,
    }


def render_pdf(doc: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    buf = io.BytesIO()
    styles = getSampleStyleSheet()

    def marking(canvas, _doc):
        canvas.saveState()
        canvas.setFont("Helvetica-Bold", 9)
        width, height = A4
        canvas.drawCentredString(width / 2, height - 12 * mm, doc["marking"])
        canvas.drawCentredString(width / 2, 10 * mm, doc["marking"])
        canvas.setFont("Helvetica", 7)
        canvas.drawString(15 * mm, 5 * mm, f"{doc['reference']} | {doc['trace']} | {doc['exported_by']}")
        canvas.drawRightString(width - 15 * mm, 5 * mm, f"Page {canvas.getPageNumber()}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(buf, pagesize=A4, topMargin=22 * mm, bottomMargin=20 * mm,
                            title=doc["title"], author=doc["author"])
    p = lambda text, style="BodyText": Paragraph(escape(text or "").replace("\n", "<br/>"), styles[style])  # noqa: E731
    story = [
        p(doc["type"], "Heading3"), p(doc["title"], "Title"),
        p(f"Référence : {doc['reference']}   Dossier : {doc['case_reference']}   Date : {doc['date']}"),
        p(f"Auteur : {doc['author']}   Validé par : {doc['validator']}"),
        p(f"Diffusion : {doc['dissemination']}"), Spacer(1, 6 * mm),
    ]
    for heading, body in doc["sections"]:
        story += [p(heading, "Heading2"), p(body or "—"), Spacer(1, 3 * mm)]
    if doc["statements"]:
        story.append(p("Éléments référencés / Referenced items", "Heading2"))
        story += [p(f"[{ref}] {kind} : {text}") for kind, text, ref in doc["statements"]]
    story.append(p("Sources", "Heading2"))
    story += [p(line) for line in doc["sources"]] or [p("—")]
    story += [Spacer(1, 6 * mm), p(doc["notice"], "Italic"),
              p(f"Empreinte du contenu validé : {doc['content_sha256']}", "Italic")]
    pdf.build(story, onFirstPage=marking, onLaterPages=marking)
    return buf.getvalue()


def render_docx(doc: dict) -> bytes:
    from docx import Document as Docx

    d = Docx()
    section = d.sections[0]
    section.header.paragraphs[0].text = doc["marking"]
    section.footer.paragraphs[0].text = f"{doc['marking']} | {doc['reference']} | {doc['trace']} | {doc['exported_by']}"
    d.add_paragraph(doc["type"])
    d.add_heading(doc["title"], level=0)
    d.add_paragraph(f"Référence : {doc['reference']}   Dossier : {doc['case_reference']}   Date : {doc['date']}")
    d.add_paragraph(f"Auteur : {doc['author']}   Validé par : {doc['validator']}")
    d.add_paragraph(f"Diffusion : {doc['dissemination']}")
    for heading, body in doc["sections"]:
        d.add_heading(heading, level=1)
        d.add_paragraph(body or "—")
    if doc["statements"]:
        d.add_heading("Éléments référencés / Referenced items", level=1)
        for kind, text, ref in doc["statements"]:
            d.add_paragraph(f"[{ref}] {kind} : {text}", style="List Bullet")
    d.add_heading("Sources", level=1)
    for line in doc["sources"] or ["—"]:
        d.add_paragraph(line, style="List Bullet")
    d.add_paragraph().add_run(doc["notice"]).italic = True
    d.add_paragraph(f"Empreinte du contenu validé : {doc['content_sha256']}")
    d.core_properties.title = doc["title"]
    d.core_properties.author = doc["author"]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

